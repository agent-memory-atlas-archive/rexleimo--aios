import json, os, sys, re
from pathlib import Path
import urllib.request

BASE='runs/jb8/jb/dataset/management_analysts/task1'
TASKDIR='dependencies/job-bench-eval/dataset/main/management_analysts/task1'
agnes=json.load(open(os.path.expanduser('~/.pi/agent/models.json')))['providers']['agnes']

def extract(p):
    s=str(p).lower()
    try:
        if s.endswith('.docx'):
            import docx
            d=docx.Document(str(p)); parts=[x.text for x in d.paragraphs]
            for t in d.tables:
                for row in t.rows: parts.append(' | '.join(c.text for c in row.cells))
            return '\n'.join(parts)
        if s.endswith('.xlsx'):
            import openpyxl
            wb=openpyxl.load_workbook(str(p),data_only=True); out=[]
            for ws in wb.worksheets:
                out.append(f'== tab: {ws.title} ==')
                for row in ws.iter_rows(values_only=True):
                    if any(v is not None for v in row): out.append(' | '.join('' if v is None else str(v) for v in row))
            return '\n'.join(out)
        if s.endswith('.csv'):
            return Path(p).read_text(errors='replace')[:6000]
        return Path(p).read_text(errors='replace')[:4000]
    except Exception as e:
        return f'[extract error: {e}]'

def call(prompt, maxtok=4000):
    req=urllib.request.Request(agnes['baseUrl'].rstrip('/')+'/chat/completions',
      data=json.dumps({'model':'agnes-3.0-flash','messages':[{'role':'user','content':prompt}],
                       'max_tokens':maxtok,'temperature':0}).encode(),
      headers={'Content-Type':'application/json','Authorization':'Bearer '+agnes['apiKey']})
    return json.load(urllib.request.urlopen(req,timeout=300))['choices'][0]['message']['content']

rubrics=json.load(open(f'{TASKDIR}/RUBRICS.json'))['rubrics']
taskcard=Path(f'{TASKDIR}/task_card.md').read_text(errors='replace')[:3500]
arms=sys.argv[1:] or ['base1']
for arm in arms:
    d=Path(f'{BASE}/model_output/jb-oc-agg1__management_analysts__task1__{arm}')
    files=[f for f in sorted(d.glob('**/*'),key=lambda x:('output/answer' in str(x))) if f.is_file()]
    blob='\n\n'.join(f'--- {f.name} ---\n'+extract(f)[:4000] for f in files[:12])
    rl='\n'.join(f'{i+1}. [w={r["weight"]}] {r["rubric"]}' for i,r in enumerate(rubrics))
    prompt=(f'You are a strict grader for a real business task.\n\n# TASK CARD\n{taskcard}\n\n'
            f'# CANDIDATE DELIVERABLES (arm {arm})\n{blob[:16000]}\n\n'
            f'# RUBRICS (weight 1-10; answer 1.0/0.5/0.0 per rubric)\n{rl}\n\n'
            'Return ONLY a JSON array: [{"i":1,"score":0.5,"why":"short"},...] one per rubric.')
    try:
        out=call(prompt)
        j=json.loads(re.search(r'\[.*\]',out,re.S).group(0))
        tot=sum(r['weight']*float(x['score']) for r,x in zip(rubrics,j))
        mx=sum(r['weight'] for r in rubrics)
        print(f'ARM {arm}: {tot/mx*100:.1f}%  (weighted {tot:.1f}/{mx})  items={len(j)}')
        json.dump({'arm':arm,'pct':tot/mx*100,'detail':j}, open('runs/RESULT_quickjudge_%s.json' % arm, 'w'), indent=1)
    except Exception as e:
        print(f'ARM {arm}: FAIL {type(e).__name__} {str(e)[:120]}')
